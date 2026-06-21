"""
Pruebas para el generador de hojas de cálculo Excel (openpyxl).
"""

import io

import openpyxl

from excel_generator import generate_invoice_excel


def test_generate_invoice_excel_binary_signature(sample_invoice_list):
    """Verifica que el output sea un BytesIO válido con firma PK de archivo ZIP/XLSX."""
    stream = generate_invoice_excel(sample_invoice_list)
    assert isinstance(stream, io.BytesIO)
    content = stream.getvalue()
    assert content.startswith(b"PK\x03\x04")
    assert len(content) > 3000


def test_generate_invoice_excel_content(sample_invoice_list):
    """Inspecciona las celdas, valores numéricos y fórmulas generadas."""
    stream = generate_invoice_excel(sample_invoice_list)
    wb = openpyxl.load_workbook(stream)
    ws = wb.active
    assert ws.title == "Consolidado de Facturas"

    # Verificar encabezado
    assert ws.cell(row=1, column=1).value == "N°"
    assert ws.cell(row=1, column=3).value == "N° Factura"
    assert ws.cell(row=1, column=11).value == "Total"

    # Verificar primera fila de datos (Fila 2)
    assert ws.cell(row=2, column=3).value == sample_invoice_list[0].numero_factura
    assert ws.cell(row=2, column=11).value == sample_invoice_list[0].total

    # Verificar fila de totales (Fila 4: 1 header + 2 facturas + 1 total)
    total_row = 4
    assert ws.cell(row=total_row, column=1).value == "TOTALES"
    # Subtotal fórmula en columna 9 (I)
    assert "=SUM(I2:I3)" in str(ws.cell(row=total_row, column=9).value)
    # Total fórmula en columna 11 (K)
    assert "=SUM(K2:K3)" in str(ws.cell(row=total_row, column=11).value)


def test_generate_invoice_excel_empty():
    """Valida la generación de una hoja vacía sin errores si no hay facturas."""
    stream = generate_invoice_excel([])
    assert isinstance(stream, io.BytesIO)
    wb = openpyxl.load_workbook(stream)
    assert wb.active is not None
